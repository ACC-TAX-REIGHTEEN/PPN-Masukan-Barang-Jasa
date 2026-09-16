import os
import sys
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

def format_tanggal_indonesia(tanggal_obj):
    if pd.isnull(tanggal_obj):
        return ""
    
    bulan_map = {
        1: 'Jan', 2: 'Feb', 3: 'Mar', 4: 'Apr', 5: 'Mei', 6: 'Jun',
        7: 'Jul', 8: 'Agu', 9: 'Sep', 10: 'Okt', 11: 'Nop', 12: 'Des'
    }
    
    try:
        tgl = tanggal_obj.day
        bln = bulan_map[tanggal_obj.month]
        thn = tanggal_obj.year
        return f"{tgl} {bln} {thn}"
    except AttributeError:
        return str(tanggal_obj)

def auto_fit_columns(filename, sheet_name):
    wb = load_workbook(filename)
    if sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
    else:
        ws = wb.active
    
    for column in ws.columns:
        max_length = 0
        column = [cell for cell in column]
        col_letter = get_column_letter(column[0].column)
        
        for cell in column:
            try:
                if cell.value is not None:
                    length = len(str(cell.value))
                    if length > max_length:
                        max_length = length
            except:
                pass
        
        adjusted_width = (max_length + 2)
        ws.column_dimensions[col_letter].width = adjusted_width
        
    wb.save(filename)
    print(f"--> Format auto-fit selesai untuk: {filename} (Sheet: {sheet_name})")

def parse_nominal(val):
    if pd.isnull(val) or val is None:
        return None
    val_str = str(val).replace('.', '').replace(',', '.').strip()
    try:
        return float(val_str)
    except ValueError:
        return None

def baca_filter_txt_kondisional(nama_file):
    if not os.path.exists(nama_file):
        print(f"--> Peringatan: File {nama_file} tidak ditemukan.")
        return []
    
    rules = []
    with open(nama_file, 'r', encoding='utf-8') as f:
        for line in f:
            line_str = line.strip()
            if not line_str:
                continue
            
            if '|' in line_str:
                parts = line_str.split('|')
                vendor = parts[0].strip()
                nom = parse_nominal(parts[1])
                rules.append({'vendor': vendor, 'nominal': nom})
            else:
                rules.append({'vendor': line_str, 'nominal': None})
                
    return rules

def cari_nama_kolom(df, daftar_kemungkinan):
    for col in df.columns:
        for nama in daftar_kemungkinan:
            if nama.lower() in str(col).lower():
                return col
    return None

def pisahkan_jv_dan_barang(df, rules_jv, rules_brg):
    col_penjual = cari_nama_kolom(df, [
        "Nama Penjual Barang Kena Pajak/Barang Kena Pajak Tidak Berwujud/Jasa Kena Pajak",
        "Nama Penjual",
        "Nama Pemasok"
    ])
    
    col_ppn = cari_nama_kolom(df, ["PPN", "Nilai PPN", "Jumlah PPN"])

    if not col_penjual:
        print("--> Error: Kolom nama penjual tidak ditemukan di Coretax.")
        return pd.DataFrame(columns=df.columns), pd.DataFrame(columns=df.columns)

    indices_jv = []
    indices_brg = []

    for idx, row in df.iterrows():
        penjual_val = str(row[col_penjual]).upper()
        ppn_val = parse_nominal(row[col_ppn]) if col_ppn else None

        is_jv = False
        for rule in rules_jv:
            vendor_kw = rule['vendor'].upper()
            rule_nom = rule['nominal']
            
            if vendor_kw in penjual_val:
                if rule_nom is not None:
                    if ppn_val is not None and abs(ppn_val - rule_nom) < 1.0:
                        is_jv = True
                        break
                else:
                    is_jv = True
                    break

        if is_jv:
            indices_jv.append(idx)
        else:
            indices_brg.append(idx)

    df_jv = df.loc[indices_jv].copy()
    df_barang = df.loc[indices_brg].copy()

    return df_jv, df_barang

def main():
    print("--> Mulai Proses Pemisahan Data Coretax")
    file_sumber = 'Coretaxm.xlsx'
    print(f"--> Membaca {file_sumber}...")

    try:
        df = pd.read_excel(file_sumber, sheet_name='data')
    except FileNotFoundError:
        print(f"--> Error: File {file_sumber} tidak ditemukan di folder me ini.")
        sys.exit()

    col_tanggal_v1 = 'Tanggal Faktur Pajak'
    col_tanggal_v2 = 'Faktur Pajak/Dokumen Tertentu/Nota Retur/Nota Pembatalan - Tanggal'

    if col_tanggal_v1 in df.columns:
        col_tanggal = col_tanggal_v1
    elif col_tanggal_v2 in df.columns:
        col_tanggal = col_tanggal_v2
    else:
        col_tanggal = None

    if col_tanggal:
        df[col_tanggal] = pd.to_datetime(df[col_tanggal])
        df[col_tanggal] = df[col_tanggal].apply(format_tanggal_indonesia)

    rules_jv = baca_filter_txt_kondisional('hjv.txt')
    rules_brg = baca_filter_txt_kondisional('hbrg.txt')

    df_jv, df_barang = pisahkan_jv_dan_barang(df, rules_jv, rules_brg)

    files_to_save = {
        'CtxJV_temp.xlsx': (df_jv, 'CoretaxJV'),
        'CtxBarang_temp.xlsx': (df_barang, 'CoretaxBarang')
    }

    for filename, (data, sheetname) in files_to_save.items():
        print(f"--> Menyimpan file: {filename} dengan Sheet: {sheetname}...")
        data.to_excel(filename, index=False, sheet_name=sheetname)
        auto_fit_columns(filename, sheetname)

    print("--> Selesai! Pemisahan data berhasil diproses.")

if __name__ == "__main__":
    main()