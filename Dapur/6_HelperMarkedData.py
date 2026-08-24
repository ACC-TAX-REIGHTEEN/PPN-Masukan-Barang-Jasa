import pandas as pd
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment

FILE_HASIL = "Hasil_Analisis_Barang_Dan_Jasa.xlsx"
FILE_CTX = "CtxJV_temp.xlsx"
SHEET_NAME = "Analisis JV"
MODE_WARNA = "BARIS_DATA"

def clean_str(val):
    if pd.isna(val) or val is None:
        return ""
    return str(val).strip().upper()

def parse_date(val):
    if pd.isna(val) or val is None:
        return ""
    val_str = str(val).strip()
    if not val_str:
        return ""
    if len(val_str) >= 10 and val_str[4] == '-' and val_str[7] == '-':
        try:
            return pd.to_datetime(val_str[:10], format='%Y-%m-%d').strftime('%Y-%m-%d')
        except Exception:
            pass
    try:
        dt = pd.to_datetime(val_str, dayfirst=True, errors='coerce')
        if pd.notna(dt):
            return dt.strftime('%Y-%m-%d')
    except Exception:
        pass
    return val_str

def clean_ppn(val):
    if pd.isna(val) or val is None:
        return None
    if isinstance(val, (int, float)):
        return round(float(val), 2)
    val_str = str(val).strip()
    if not val_str:
        return None
    if '.' in val_str and ',' in val_str:
        if val_str.find('.') < val_str.find(','):
            val_str = val_str.replace('.', '').replace(',', '.')
        else:
            val_str = val_str.replace(',', '')
    elif '.' in val_str:
        parts = val_str.split('.')
        if len(parts) > 1 and len(parts[-1]) == 3:
            val_str = val_str.replace('.', '')
    elif ',' in val_str:
        parts = val_str.split(',')
        if len(parts) > 1 and len(parts[-1]) == 2:
            val_str = val_str.replace(',', '.')
        else:
            val_str = val_str.replace(',', '')
    try:
        return round(float(val_str), 2)
    except Exception:
        return None

def get_cell_value(ws, row, col):
    cell = ws.cell(row=row, column=col)
    if cell.value is not None:
        return cell.value
    for rng in ws.merged_cells.ranges:
        if cell.coordinate in rng:
            return ws.cell(row=rng.min_row, column=rng.min_col).value
    return None

def jalankan_rekonsiliasi():
    print("--> 1. Membaca data referensi Coretax...")
    df_ctx = pd.read_excel(FILE_CTX)
    ctx_dict = {}
    for idx, row in df_ctx.iterrows():
        tgl = parse_date(row.get("Tanggal Faktur Pajak"))
        nama = clean_str(row.get("Nama Penjual"))
        ppn = clean_ppn(row.get("PPN"))
        status = str(row.get("Status Faktur")).strip() if pd.notna(row.get("Status Faktur")) else ""
        if tgl and nama and ppn is not None:
            ctx_dict[(tgl, nama, ppn)] = status
    print(f"--> 2. Membuka file {FILE_HASIL} (Sheet: {SHEET_NAME})...")
    wb = openpyxl.load_workbook(FILE_HASIL)
    if SHEET_NAME not in wb.sheetnames:
        raise ValueError(f"Sheet '{SHEET_NAME}' tidak ditemukan dalam file!")
    ws = wb[SHEET_NAME]
    target_col = 8
    header_cell = ws.cell(row=1, column=target_col, value="Status Faktur")
    header_cell.font = Font(bold=False, size=11)
    header_cell.alignment = Alignment(horizontal="center", vertical="center")
    pink_fill = PatternFill(start_color="F8CECC", end_color="F8CECC", fill_type="solid")
    yellow_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    print("--> 3. Memproses pencocokan dan pemberian warna...")
    for r in range(2, ws.max_row + 1):
        val_a = get_cell_value(ws, r, 1)
        if val_a and "TOTAL" in str(val_a).upper():
            continue
        tgl_raw = get_cell_value(ws, r, 1)
        nama_raw = get_cell_value(ws, r, 2)
        ppn_raw = get_cell_value(ws, r, 3)
        if tgl_raw is None and r - 1 >= 2:
            tgl_raw = get_cell_value(ws, r - 1, 1)
        if nama_raw is None and r - 1 >= 2:
            nama_raw = get_cell_value(ws, r - 1, 2)
        if ppn_raw is None:
            if r + 1 <= ws.max_row and ws.cell(row=r+1, column=3).value is not None:
                ppn_raw = ws.cell(row=r+1, column=3).value
            elif r - 1 >= 2 and ws.cell(row=r-1, column=3).value is not None:
                ppn_raw = ws.cell(row=r-1, column=3).value
        tgl = parse_date(tgl_raw)
        nama = clean_str(nama_raw)
        ppn = clean_ppn(ppn_raw)
        key = (tgl, nama, ppn)
        if key in ctx_dict:
            status_val = ctx_dict[key]
            cell_status = ws.cell(row=r, column=target_col, value=status_val)
            cell_status.alignment = Alignment(horizontal="center", vertical="center")
            status_upper = status_val.upper()
            fill_to_apply = None
            if "CANCELED" in status_upper or "BATAL" in status_upper:
                fill_to_apply = pink_fill
            elif "PENDING" in status_upper:
                fill_to_apply = yellow_fill
            if fill_to_apply:
                if MODE_WARNA == "HANYA_STATUS":
                    cell_status.fill = fill_to_apply
                elif MODE_WARNA == "BARIS_DATA":
                    for c in range(1, target_col + 1):
                        ws.cell(row=r, column=c).fill = fill_to_apply
    wb.save(FILE_HASIL)
    print(f"--> 4. Selesai! File {FILE_HASIL} berhasil diperbarui.")

if __name__ == "__main__":
    jalankan_rekonsiliasi()
