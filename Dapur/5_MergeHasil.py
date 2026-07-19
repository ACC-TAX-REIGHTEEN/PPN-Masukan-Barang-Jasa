from copy import copy
import sys
import openpyxl

def main():
    try:
        wb_jv = openpyxl.load_workbook("Hasil_Analisis_JV_temp.xlsx")
        wb_barang = openpyxl.load_workbook("Hasil_Analisis_Barang_temp.xlsx")
    except Exception as e:
        print(f"--> Gagal membuka file: {e}")
        sys.exit()

    ws_src = wb_jv["Analisis"]
    ws_dest = wb_barang.create_sheet(title="Analisis JV")
    
    for row in ws_src.rows:
        for cell in row:
            new_cell = ws_dest.cell(
                row=cell.row, column=cell.column, value=cell.value
            )
            if cell.has_style:
                new_cell.font = copy(cell.font)
                new_cell.border = copy(cell.border)
                new_cell.fill = copy(cell.fill)
                new_cell.number_format = copy(cell.number_format)
                new_cell.protection = copy(cell.protection)
                new_cell.alignment = copy(cell.alignment)

        for merged_cell in ws_src.merged_cells.ranges:
            ws_dest.merge_cells(str(merged_cell))

    for col_name, col_dim in ws_src.column_dimensions.items():
        ws_dest.column_dimensions[col_name].width = col_dim.width

    for row_name, row_dim in ws_src.row_dimensions.items():
        ws_dest.row_dimensions[row_name].height = row_dim.height
        
    try:
        if "Summary Total" in wb_barang.sheetnames:
            ws_summary = wb_barang["Summary Total"]
            
            last_row_jv = ws_dest.max_row
            ppn_coretax_val = ws_dest.cell(row=last_row_jv, column=3).value 
            ppn_accurate_val = ws_dest.cell(row=last_row_jv, column=5).value
            
            total_row = None
            for r in range(1, ws_summary.max_row + 1):
                if ws_summary.cell(row=r, column=1).value == "TOTAL":
                    total_row = r
                    break
                    
            if total_row is None:
                total_row = ws_summary.max_row
                
            jv_row = total_row + 1
            ws_summary.cell(row=jv_row, column=1, value="JV")
            ws_summary.cell(row=jv_row, column=2, value=ppn_coretax_val)
            ws_summary.cell(row=jv_row, column=3, value=ppn_accurate_val)
            
            sum_row = total_row + 2
            ws_summary.cell(row=sum_row, column=2, value=f"=B{total_row}+B{jv_row}") 
            ws_summary.cell(row=sum_row, column=3, value=f"=C{total_row}+C{jv_row}")
            ws_summary.cell(row=sum_row, column=4, value=f"=B{sum_row}-C{sum_row}") 
            
            fmt = ws_summary.cell(row=total_row, column=2).number_format
            ws_summary.cell(row=jv_row, column=2).number_format = fmt
            ws_summary.cell(row=jv_row, column=3).number_format = fmt
            ws_summary.cell(row=sum_row, column=2).number_format = fmt
            ws_summary.cell(row=sum_row, column=3).number_format = fmt
            ws_summary.cell(row=sum_row, column=4).number_format = fmt
            
            print("--> Berhasil menambahkan data dan rumus ke Summary Total.")
        else:
            print("--> Peringatan: Sheet 'Summary Total' tidak ditemukan di file barang.")
            
    except Exception as e:
        print(f"--> Gagal memproses Summary Total: {e}")

    try:
        wb_barang.save("Hasil_Analisis_Barang_Dan_Jasa.xlsx")
        print("--> Berhasil menggabungkan file dan menyimpannya...")
    except Exception as e:
        print(f"--> Gagal menyimpan file: {e}")
    finally:
        wb_jv.close()
        wb_barang.close()

if __name__ == "__main__":
    main()